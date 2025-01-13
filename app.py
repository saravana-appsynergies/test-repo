import streamlit as st
from docx import Document

st.title("🎉 Dynamic DOCX Generator 🎉")
st.markdown("""
Welcome to the **Dynamic DOCX Generator**! 🌟  
Fill in the details below to generate a personalized DOCX document.  
You can customize the name, location, and company in your document template.
""")


# Function to generate docx file with dynamic values
def generate_docx(name, location, company):
    # Load the existing template
    doc = Document("template.docx")  # Make sure the template.docx is in the same directory

    # Loop through each paragraph and replace placeholders
    for para in doc.paragraphs:
        if '<Name>' in para.text:
            para.text = para.text.replace('<Name>', name)
        if '<Location>' in para.text:
            para.text = para.text.replace('<Location>', location)
        if '<Company>' in para.text:
            para.text = para.text.replace('<Company>', company)

    # Save the modified document
    file_path = "generated_document.docx"
    doc.save(file_path)

    return file_path

# Streamlit inputs
name = st.text_input('📝 Enter your name:')
location = st.text_input('🌍 Enter the location:')
company = st.text_input('🏢 Enter the company name:')

# Validation and DOCX generation
if st.button('Generate DOCX ✨'):
    if name and location and company:
        docx_file = generate_docx(name, location, company)
        st.success('✅ DOCX file generated successfully!')

        # Provide download link for the generated DOCX file
        with open(docx_file, "rb") as file:
            st.download_button("🔽 Download DOCX", file, file_name="generated_document.docx")
    else:
        st.error("⚠️ Please fill out all fields!")
